"""Phase 0A skeleton for idea-to-plan (10 slides, 1920x1080).

Spec: %USERPROFILE%\\Downloads\\Canva-PPT-Prompts\\set9-newyear-business-gamenight\\
      prompt-set9-4-business-plan-idea-to-plan.md
Concept "Aurora Planner": Orchid Mist pages, aurora blobs, STEP pill on slides 2-10,
Idea Bulb buddy, frosted sticker cards with number badges, gradient squiggle under titles.
Fonts: Archivo Black (titles), Poppins (body; bold flag for labels).
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "idea-to-plan-skeleton.pptx")
PLACEHOLDERS = []  # (slide, label, x, y, w, h)


def px(v):
    return Emu(int(round(v * 9525)))


def fpt(pxsize):
    return Pt(pxsize * 0.75)


C = dict(
    mist="FBF5FF", ink="2B1A4A", orchid="7A45D6", coral="FF7A66", bubble="FF8FC7",
    peri="7C8BFF", mint="5FE0C4", lemon="FFD66B", white="FFFFFF",
    lilac="CDBBF2", shadow="E9DFFA",
)


def rgb(h):
    return RGBColor.from_string(h)


def tint(h, pct=0.2):
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    f = lambda c: int(round(c * pct + 255 * (1 - pct)))
    return "%02X%02X%02X" % (f(r), f(g), f(b))


def on(color):
    """White text only on Deep Orchid; Ink Plum everywhere else."""
    return C["white"] if color == C["orchid"] else C["ink"]


TITLE_F = "Archivo Black"
BODY_F = "Poppins"
BOLD_F = "Poppins"  # bold via bold flag

prs = Presentation()
prs.slide_width = Emu(18288000)
prs.slide_height = Emu(10287000)
BLANK = prs.slide_layouts[6]


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(C["mist"])


def _alpha(sp, alpha):
    clr = sp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=0, radius=None, alpha=None,
          name=None, rot=0):
    sp = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if name:
        sp.name = name
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _alpha(sp, alpha)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = px(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.5, radius / min(w, h))
    if rot:
        sp.rotation = rot
    if sp.has_text_frame:
        sp.text_frame.text = ""
    return sp


def rrect(slide, x, y, w, h, fill, radius=28, **kw):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, radius=radius, **kw)


def rect(slide, x, y, w, h, fill, **kw):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, **kw)


def circle(slide, x, y, d, fill, **kw):
    return shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill, **kw)


def dotted(slide, x1, y1, x2, y2, color, wpx=3, dash=MSO_LINE.ROUND_DOT, name=None):
    """Plain dotted connector line."""
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = px(wpx)
    ln.line.dash_style = dash
    if name:
        ln.name = name
    return ln


def text(slide, x, y, w, h, runs, size, color, font=BODY_F, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line_spacing=None, name=None, rot=0):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, font, bold[, size])."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    if name:
        tb.name = name
    if rot:
        tb.rotation = rot
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        segs = para if isinstance(para, list) else [(para, font, bold)]
        for seg in segs:
            t, f, b = seg[0], seg[1], seg[2]
            sz = seg[3] if len(seg) > 3 else size
            r = p.add_run()
            r.text = t
            r.font.name = f
            r.font.size = fpt(sz)
            r.font.bold = b
            r.font.color.rgb = rgb(color)
            if spacing is not None:
                r._r.get_or_add_rPr().set("spc", str(int(spacing)))
    return tb


def ph(slide, idx, label, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE):
    """Labeled placeholder rectangle, fill 20% transparent, dashed outline."""
    sp = shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill="E6D6F0", alpha=0.8,
               line=C["ink"], line_w=2, name=label)
    sp.line.dash_style = MSO_LINE.DASH
    tf = sp.text_frame
    tf.word_wrap = True
    small = min(w, h) <= 80
    tf.margin_left = tf.margin_right = px(2 if small else 8)
    tf.margin_top = tf.margin_bottom = px(2 if small else 4)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = BOLD_F
    r.font.bold = True
    r.font.size = fpt(12 if small else 16)
    r.font.color.rgb = rgb(C["ink"])
    PLACEHOLDERS.append((idx, label, x, y, w, h))
    return sp


def eyebrow(slide, x, y, t, color, align=PP_ALIGN.LEFT):
    """Small tinted chip + Poppins Bold 22 UPPERCASE, letter spacing 200 (spc 330)."""
    cw = int(len(t) * 18 + 56)
    if align == PP_ALIGN.CENTER:
        x = (1920 - cw) // 2
    rrect(slide, x, y, cw, 40, tint(color, 0.45), radius=20)
    return text(slide, x, y, cw, 40, t, 22, C["ink"], font=BOLD_F, bold=True, spacing=330,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def title(slide, x, y, w, h, t, size=84, align=PP_ALIGN.LEFT):
    return text(slide, x, y, w, h, t, size, C["ink"], font=TITLE_F, align=align,
                anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)


SQ = "DECOR: squiggle underline (Orchid>Bubblegum>Coral)"


def squiggle(slide, idx, x, y, w=420):
    ph(slide, idx, SQ, x, y, w, 26)


def step_pill(slide, n):
    """Deep Orchid pill top right: 'STEP nn / 10' + small arrow circle; doubles as page number."""
    x, y, w, h = 1560, 56, 290, 52
    rrect(slide, x, y, w, h, C["orchid"], radius=26)
    text(slide, x + 28, y, 200, h, "STEP %02d / 10" % n, 22, C["white"], font=BOLD_F, bold=True,
         spacing=198, anchor=MSO_ANCHOR.MIDDLE)
    circle(slide, x + w - 44, y + 8, 36, C["white"])
    shape(slide, MSO_SHAPE.RIGHT_ARROW, x + w - 37, y + 18, 22, 16, fill=C["orchid"])


def badge(slide, x, y, n, col, d=64):
    circle(slide, x, y, d, col, line=C["white"], line_w=4)
    text(slide, x, y, d, d, str(n), 28, on(col), font=BOLD_F, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)


def scard(slide, x, y, w, h, n=None, col=None, shadow=True):
    """Frosted sticker card: white, radius 32, 3px lilac outline, soft shadow, corner badge."""
    if shadow:
        rrect(slide, x + 4, y + 8, w, h, C["shadow"], radius=32)
    rrect(slide, x, y, w, h, C["white"], radius=32, line=C["lilac"], line_w=3)
    if n is not None:
        badge(slide, x - 14, y - 14, n, col)


def card_side(slide, idx, x, y, w, h, n, col, label, body, icon, isz=64):
    """Sticker card: icon left, bold label + body text right."""
    scard(slide, x, y, w, h, n, col)
    ph(slide, idx, "ICON: " + icon, x + 30, y + 56, isz, isz)
    tx = x + 30 + isz + 24
    text(slide, tx, y + 18, w - (tx - x) - 24, 44, label, 28, C["ink"], font=BOLD_F, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, tx, y + 66, w - (tx - x) - 24, h - 66 - 16, body, 26, C["ink"],
         line_spacing=1.0)


def pill(slide, x, y, w, h, color, t, size=28, icon=None, idx=None, tcolor=None, line=None,
         bold=True, fill_tint=None):
    rrect(slide, x, y, w, h, fill_tint or color, radius=h // 2, line=line,
          line_w=3 if line else 0)
    tx = x
    if icon:
        isz = 56 if h >= 96 else 44
        ph(slide, idx, "ICON: " + icon, x + 18, y + (h - isz) // 2, isz, isz)
        tx = x + 18 + isz + 12
    text(slide, tx, y, w - (tx - x) - (14 if icon else 0), h, t, size,
         tcolor or on(color), font=BOLD_F, bold=bold,
         align=PP_ALIGN.LEFT if icon else PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def blob(slide, idx, x, y, w, h, mix="Orchid > Bubblegum > Coral", anchor=MSO_ANCHOR.MIDDLE):
    ph(slide, idx, "DECOR: aurora blob, soft blurred gradient (%s)" % mix, x, y, w, h,
       anchor=anchor)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


NOTE_GEN = ("General information only. Every name and number in this deck is made up for "
            "illustration. Replace them with your own research.")


def new():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    return s


# ---------------- Slide 1: Cover (title left, hero right; no STEP pill) ----------------
s = new()
ph(s, 1, "DECOR: very large aurora blob bleeding off top-right corner (Orchid > Bubblegum > "
         "Coral)", 1250, 0, 670, 540, anchor=MSO_ANCHOR.TOP)
ph(s, 1, "DECOR: aurora blob bleeding off bottom-left (Periwinkle > Aqua)", 0, 900, 640, 180)
ph(s, 1, "DECOR: sparkle confetti", 100, 30, 380, 90)
eyebrow(s, 120, 150, "A CREATIVE FOUNDER'S GUIDE", C["lemon"])
title(s, 120, 215, 940, 340, ["Turn Your Idea", "Into a", "Business Plan"], size=96)
ph(s, 1, SQ, 120, 566, 520, 30)
pill(s, 120, 628, 600, 76, C["orchid"], "Dream it. Plan it. Launch it.")
text(s, 120, 730, 900, 40, "Presented by [Your Name] · [Organization]", 26, C["ink"],
     anchor=MSO_ANCHOR.MIDDLE)
for i, (col, ic) in enumerate([(C["bubble"], "lightbulb"), (C["mint"], "notebook"),
                               (C["lemon"], "rocket")]):
    cx = 120 + i * 130
    circle(s, cx, 790, 104, tint(col, 0.55), line=col, line_w=3)
    ph(s, 1, "ICON: " + ic, cx + 22, 812, 60, 60)
ph(s, 1, "HERO: cheerful maker at a worktable w/ notebook, flower buckets + coffee cups, Idea "
         "Bulb buddy floating above, paper rocket flying out 780w", 1060, 150, 780, 760)
ph(s, 1, "DECOR: sparkles + hearts", 1000, 940, 400, 100)

# ---------------- Slide 2: Why Write a Plan? (center hero, 4 callouts) ----------------
s = new()
step_pill(s, 2)
ph(s, 2, "DECOR: aurora blob TL (Orchid > Bubblegum)", 0, 0, 330, 240)
ph(s, 2, "DECOR: aurora blob BL (Periwinkle > Aqua)", 0, 860, 310, 220)
ph(s, 2, "DECOR: aurora blob BR (Bubblegum > Coral)", 1610, 860, 310, 220)
ph(s, 2, "DECOR: sparkle confetti", 1480, 120, 360, 110)
eyebrow(s, 0, 44, "WHY IT MATTERS", C["peri"], align=PP_ALIGN.CENTER)
title(s, 360, 96, 1200, 110, "Why Write a Plan?", align=PP_ALIGN.CENTER)
ph(s, 2, SQ, 760, 208, 400, 26)
ph(s, 2, "HERO: friendly open notebook w/ roadmap + flag drawn on it, Idea Bulb buddy sitting on "
         "top 520w", 700, 280, 520, 540)
cards2 = [
    (100, 280, 1, C["orchid"], "Get Clear", "target",
     "Turn a fuzzy idea into clear goals you can act on."),
    (100, 560, 2, C["coral"], "Spot Risks", "shield check",
     "Find problems on paper before they cost real money."),
    (1300, 280, 3, C["peri"], "Win Support", "handshake",
     "Banks, partners and investors want to see your thinking."),
    (1300, 560, 4, C["mint"], "Stay on Track", "calendar",
     "Use it as a map to check your progress every month."),
]
for cx, cy, n, col, lab, ic, body in cards2:
    card_side(s, 2, cx, cy, 520, 250, n, col, lab, body, ic)
    if cx < 700:
        dotted(s, 622, cy + 125, 698, cy + 125, C["orchid"])
    else:
        dotted(s, 1222, cy + 125, 1298, cy + 125, C["orchid"])
rrect(s, 330, 856, 1260, 84, tint(C["lemon"], 0.5), radius=42)
text(s, 360, 856, 1200, 84, "A plan can be one page or fifty. Short and honest beats long and vague.",
     26, C["ink"], font=BOLD_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, NOTE_GEN)

# ---------------- Slide 3: The Executive Summary (4 bubbles left, hero right) ----------------
s = new()
step_pill(s, 3)
ph(s, 3, "DECOR: large aurora blob behind hero, bleeding off right edge (Orchid > Bubblegum > "
         "Coral)", 1260, 250, 660, 720, anchor=MSO_ANCHOR.TOP)
ph(s, 3, "DECOR: aurora blob BL (Periwinkle > Aqua)", 0, 980, 420, 100)
ph(s, 3, "DECOR: sparkle confetti", 1500, 120, 360, 100)
eyebrow(s, 100, 44, "SECTION 1", C["bubble"])
title(s, 100, 96, 1400, 110, "The Executive Summary")
squiggle(s, 3, 100, 208, 420)
text(s, 100, 254, 1150, 44, "Write it last, put it first. One page that tells the whole story.",
     26, C["ink"], anchor=MSO_ANCHOR.MIDDLE)
sent3 = [
    (C["orchid"], "We are [Your Business Name], a [type of business] in [location]."),
    (C["coral"], "We help [your customers] solve [their problem]."),
    (C["peri"], "We make money by [how you earn]."),
    (C["mint"], "We need [amount] to [what the money will do]."),
]
for i, (col, t) in enumerate(sent3):
    by = 330 + i * 136
    scard(s, 100, by, 1120, 118, i + 1, col)
    text(s, 140, by + 8, 1040, 102, t, 26, C["ink"], anchor=MSO_ANCHOR.MIDDLE,
         line_spacing=1.0)
ph(s, 3, "HERO: person writing on a giant sheet of paper, Idea Bulb buddy holding a pencil 520w",
   1300, 300, 520, 560)
rrect(s, 100, 888, 1120, 76, tint(C["mint"], 0.45), radius=38)
text(s, 130, 888, 1060, 76, "Keep it to one page. If a stranger can repeat it, it works.", 26,
     C["ink"], font=BOLD_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, NOTE_GEN)

# ---------------- Slide 4: Problem and Solution (two big panels with an arrow) ----------------
s = new()
step_pill(s, 4)
ph(s, 4, "DECOR: aurora blob BL (Periwinkle > Aqua)", 0, 920, 420, 160)
ph(s, 4, "DECOR: aurora blob BR (Bubblegum > Coral)", 1500, 920, 420, 160)
ph(s, 4, "DECOR: sparkle confetti", 1480, 130, 380, 100)
ph(s, 4, "ICON: Idea Bulb buddy (thinking pose)", 110, 100, 130, 140)
eyebrow(s, 0, 44, "SECTION 2", C["coral"], align=PP_ALIGN.CENTER)
title(s, 360, 96, 1200, 110, "Problem and Solution", align=PP_ALIGN.CENTER)
ph(s, 4, SQ, 760, 208, 400, 26)
panels4 = [
    (100, 1, C["coral"], "The Problem",
     "Describe the pain your customers feel, in their words. Add one real-life moment.",
     "HERO: rushed person with a wilted flower and a clock 480w",
     "Example only: busy people can't find fresh flowers after work."),
    (1040, 2, C["mint"], "The Solution",
     "Show how your product or service fixes it, simply. Say why it is better than what "
     "exists.",
     "HERO: happy person with a bouquet and a coffee at a little stand 480w",
     "Example only: a late-night flower-and-coffee stand near the train station."),
]
for x0, n, col, head, body, hero, ex in panels4:
    y0, pw, pnh = 262, 780, 610
    scard(s, x0, y0, pw, pnh)
    rrect(s, x0, y0, pw, 72, col, radius=32)
    rect(s, x0, y0 + 36, pw, 36, col)
    badge(s, x0 - 14, y0 - 14, n, col)
    text(s, x0 + 56, y0, pw - 112, 72, head, 28, C["ink"], font=BOLD_F, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0 + 36, y0 + 90, pw - 72, 100, body, 26, C["ink"], line_spacing=1.0)
    ph(s, 4, hero, x0 + 150, y0 + 198, 480, 290)
    rrect(s, x0 + 36, y0 + 502, pw - 72, 94, C["white"], radius=28, line=col, line_w=3)
    text(s, x0 + 62, y0 + 502, pw - 124, 94, ex, 26, C["ink"], font=BOLD_F, bold=True,
         anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
circle(s, 912, 519, 96, C["orchid"], line=C["white"], line_w=4)
shape(s, MSO_SHAPE.RIGHT_ARROW, 934, 548, 52, 38, fill=C["white"])
pill(s, 530, 892, 860, 76, C["lemon"], "Ask yourself: who hurts, how much, and why now?")
notes(s, NOTE_GEN)

# ---------------- Slide 5: Know Your Customers (ring diagram left, 3 cards right) --------------
s = new()
step_pill(s, 5)
ph(s, 5, "DECOR: aurora blob BL (Periwinkle > Aqua)", 0, 920, 640, 160)
ph(s, 5, "DECOR: aurora blob BR (Orchid > Bubblegum > Coral)", 1480, 940, 440, 140)
ph(s, 5, "DECOR: sparkle confetti", 1520, 120, 340, 100)
ph(s, 5, "ICON: Idea Bulb buddy (pointing at rings)", 1380, 110, 120, 130)
eyebrow(s, 100, 44, "SECTION 3 · MARKET", C["peri"])
title(s, 100, 96, 1200, 110, "Know Your Customers")
squiggle(s, 5, 100, 208, 420)
# concentric rings (plain circles)
circle(s, 110, 290, 600, tint(C["peri"], 0.5), line=C["peri"], line_w=3, name="ring outer")
circle(s, 210, 390, 400, tint(C["bubble"], 0.55), line=C["bubble"], line_w=3, name="ring middle")
circle(s, 310, 490, 200, tint(C["coral"], 0.7), line=C["coral"], line_w=3, name="ring core")
for (lab, py0, pw) in [(["Everyone who", "might buy"], 246, 270),
                       (["People near you", "who need it"], 346, 300),
                       (["Your first", "100 customers"], 446, 270)]:
    rrect(s, 410 - pw // 2, py0, pw, 88, C["white"], radius=32, line=C["orchid"], line_w=3)
    text(s, 410 - pw // 2 + 10, py0, pw - 20, 88, lab, 26, C["ink"], font=BOLD_F, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
ph(s, 5, "ICON: map pin", 142, 612, 56, 56)
ph(s, 5, "ICON: people dot", 624, 574, 52, 52)
ph(s, 5, "ICON: map pin", 532, 772, 56, 56)
ph(s, 5, "ICON: people dot", 244, 794, 52, 52)
ph(s, 5, "ICON: people dot", 304, 674, 52, 52)
ph(s, 5, "ICON: map pin", 474, 664, 56, 56)
cards5 = [
    (C["orchid"], "Who are they?", "customer",
     "Age, job, habits, budget. Describe one typical customer."),
    (C["coral"], "Where are they?", "location pin",
     "Which neighborhood, website or app do they use?"),
    (C["peri"], "How big is it?", "chart",
     "Estimate how many people could buy, and how often."),
]
for i, (col, lab, ic, body) in enumerate(cards5):
    card_side(s, 5, 780, 282 + i * 184, 1040, 160, i + 1, col, lab, body, ic)
pill(s, 780, 846, 1040, 76, C["lemon"], "Start small. Win your first 100 customers, then grow.")
notes(s, NOTE_GEN)

# ---------------- Slide 6: Stand Out From the Crowd (2x2 matrix center, side cards) ------------
s = new()
step_pill(s, 6)
ph(s, 6, "DECOR: aurora blob BL (Periwinkle > Aqua)", 0, 860, 310, 220)
ph(s, 6, "DECOR: aurora blob BR (Orchid > Bubblegum > Coral)", 1610, 860, 310, 220)
eyebrow(s, 100, 44, "SECTION 4 · COMPETITION", C["mint"])
title(s, 100, 96, 1450, 110, "Stand Out From the Crowd")
squiggle(s, 6, 100, 208, 420)
# matrix card (plain shapes)
scard(s, 640, 262, 640, 580)
rect(s, 993, 292, 4, 470, C["lilac"], name="matrix vertical axis")
rect(s, 748, 525, 494, 4, C["lilac"], name="matrix horizontal axis")
text(s, 748, 778, 240, 34, "Lower price", 22, C["ink"], font=BOLD_F, bold=True,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 1010, 778, 240, 34, "Higher price", 22, C["ink"], font=BOLD_F, bold=True,
     align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
text(s, 700 - 130, 400 - 17, 260, 34, "Special experience", 22, C["ink"], font=BOLD_F, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rot=270)
text(s, 700 - 130, 650 - 17, 260, 34, "Basic experience", 22, C["ink"], font=BOLD_F, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rot=270)
rrect(s, 1050, 296, 190, 36, C["white"], radius=18, line=C["lilac"], line_w=3)
text(s, 1050, 296, 190, 36, "Example only", 22, C["ink"], font=BOLD_F, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
pill(s, 768, 372, 214, 84, C["coral"], "You", icon="star", idx=6)
for lab, col, px0, py0 in [("Competitor A", C["peri"], 752, 610),
                           ("Competitor B", C["bubble"], 1010, 345),
                           ("Competitor C", C["lemon"], 1010, 610)]:
    pill(s, px0, py0, 232, 60, col, lab, size=26)
card_side(s, 6, 100, 330, 500, 260, 1, C["orchid"], "List your rivals",
          "Name 3–5 businesses your customers could choose instead.", "checklist")
card_side(s, 6, 1320, 330, 500, 260, 2, C["mint"], "Find your edge",
          "What do you do better, faster, cheaper or kinder?", "shield check")
ph(s, 6, "ICON: Idea Bulb buddy (holding magnifier)", 110, 640, 190, 200)
ph(s, 6, "DECOR: sparkle confetti", 1330, 640, 440, 190)
rrect(s, 330, 868, 1260, 80, tint(C["lemon"], 0.5), radius=40)
text(s, 360, 868, 1200, 80, "Your edge is the reason a customer picks you on a busy day.", 26,
     C["ink"], font=BOLD_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, NOTE_GEN + "\nThe matrix is an example only. Pick two things your customers really "
      "care about, such as price and experience, and place real competitors honestly.")

# ---------------- Slide 7: Marketing and Sales (funnel left, hero right) ----------------
s = new()
step_pill(s, 7)
ph(s, 7, "DECOR: large aurora blob behind hero, bleeding off right edge (Orchid > Bubblegum > "
         "Coral)", 1330, 240, 590, 720, anchor=MSO_ANCHOR.TOP)
ph(s, 7, "DECOR: aurora blob BL (Periwinkle > Aqua)", 0, 940, 520, 140)
ph(s, 7, "DECOR: sparkle confetti", 1500, 118, 360, 90)
eyebrow(s, 100, 44, "SECTION 5", C["bubble"])
title(s, 100, 96, 1200, 110, "Marketing and Sales")
squiggle(s, 7, 100, 208, 420)
funnel = [
    (1020, C["peri"], "Notice", "They find you: posts, signs, word of mouth.", "eye"),
    (900, C["bubble"], "Care", "They like what they see: photos, reviews, a clear offer.", "heart"),
    (780, C["coral"], "Buy", "Make buying easy: simple prices, easy payment.", "shopping bag"),
    (660, C["mint"], "Return", "They come back: friendly service, small rewards.", "repeat"),
]
for i, (fw, col, lab, body, ic) in enumerate(funnel):
    fx, fy = 610 - fw // 2, 270 + i * 134
    rrect(s, fx, fy, fw, 120, col, radius=40, line=C["white"], line_w=3)
    badge(s, fx - 14, fy - 14, i + 1, C["orchid"], d=56)
    ph(s, 7, "ICON: " + ic, fx + 28, fy + 32, 56, 56)
    text(s, fx + 104, fy + 8, fw - 128, 40, lab, 28, C["ink"], font=BOLD_F, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, fx + 104, fy + 48, fw - 128, 68, body, 26, C["ink"], line_spacing=1.0)
chx = 100
for lab, w, ic in [("Social media", 300, "phone"), ("Local events", 300, "calendar"),
                   ("Email list", 266, "mail"), ("Word of mouth", 316, "chat bubble")]:
    pill(s, chx, 830, w, 64, C["lilac"], lab, size=26, icon=ic, idx=7,
         fill_tint=tint(C["orchid"], 0.18))
    chx += w + 14
ph(s, 7, "HERO: maker holding a megaphone + a phone showing a post, Idea Bulb buddy cheering 520w",
   1340, 262, 520, 640)
notes(s, NOTE_GEN)

# ---------------- Slide 8: Operations and Team (hero left, 2x2 cards, team row) ----------------
s = new()
step_pill(s, 8)
ph(s, 8, "DECOR: aurora blob behind hero (Orchid > Bubblegum > Coral)", 0, 236, 680, 470,
   anchor=MSO_ANCHOR.TOP)
ph(s, 8, "DECOR: aurora blob BL (Periwinkle > Aqua)", 0, 950, 420, 130)
ph(s, 8, "DECOR: sparkle confetti", 1470, 850, 350, 120)
ph(s, 8, "ICON: Idea Bulb buddy (holding clipboard)", 1290, 846, 140, 150)
eyebrow(s, 100, 44, "SECTIONS 6 AND 7", C["coral"])
title(s, 100, 96, 1300, 110, "Operations and Team")
squiggle(s, 8, 100, 208, 420)
ph(s, 8, "HERO: little flat market stall full of flowers + coffee, maker serving a customer 520w",
   100, 262, 520, 420)
cards8 = [
    (680, 262, 1, C["orchid"], "Location", "location pin", "Where will you work?"),
    (1270, 262, 2, C["coral"], "Supplies", "box", "What do you buy, and from whom?"),
    (680, 462, 3, C["peri"], "Daily routine", "clock", "What happens each day?"),
    (1270, 462, 4, C["mint"], "Tools", "wrench", "The equipment and software you need."),
]
for cx, cy, n, col, lab, ic, body in cards8:
    card_side(s, 8, cx, cy, 550, 180, n, col, lab, body, ic)
for i, (name, role) in enumerate([("[Founder Name]", "Founder"), ("[Team Member]", "Operations"),
                                  ("[Advisor Name]", "Advisor")]):
    tx0 = 100 + i * 585
    rrect(s, tx0, 716, 550, 110, C["lemon"], radius=55)
    ph(s, 8, "ICON: avatar", tx0 + 14, 727, 88, 88)
    text(s, tx0 + 122, 726, 550 - 122 - 30, 44, name, 28, C["ink"], font=BOLD_F, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, tx0 + 122, 770, 550 - 122 - 30, 38, role, 26, C["ink"], anchor=MSO_ANCHOR.MIDDLE)
rrect(s, 100, 858, 1100, 76, tint(C["bubble"], 0.45), radius=38)
text(s, 130, 858, 1040, 76, "No team yet? List the skills you still need to find.", 26,
     C["ink"], font=BOLD_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, NOTE_GEN)

# ---------------- Slide 9: Your Money Plan (hero left, 3 formula cards right) ----------------
s = new()
step_pill(s, 9)
ph(s, 9, "DECOR: aurora blob behind hero (Orchid > Bubblegum > Coral)", 0, 236, 640, 520,
   anchor=MSO_ANCHOR.TOP)
ph(s, 9, "DECOR: aurora blob BR (Periwinkle > Aqua)", 1500, 880, 420, 200)
ph(s, 9, "DECOR: sparkle confetti", 1500, 118, 360, 90)
eyebrow(s, 100, 44, "SECTION 8 · FINANCES", C["lemon"])
title(s, 100, 96, 1000, 110, "Your Money Plan")
squiggle(s, 9, 100, 208, 420)
ph(s, 9, "HERO: big piggy bank w/ coins, Idea Bulb buddy holding a calculator 500w",
   100, 262, 500, 480)
scard(s, 100, 770, 500, 190)
ph(s, 9, "DECOR: simple bar chart (flat, decor only)", 124, 786, 452, 100)
text(s, 124, 904, 452, 36, "Illustrative — not real data", 22, C["ink"], font=BOLD_F, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
cards9 = [
    (C["orchid"], "Startup costs", "coin stack",
     "One-time spending to open: equipment, license, first stock."),
    (C["coral"], "Monthly costs", "calendar",
     "Rent, supplies, pay, marketing and other bills every month."),
    (C["mint"], "Break-even", "scale",
     "The sales you need each month to cover all your costs."),
]
for i, (col, lab, ic, body) in enumerate(cards9):
    card_side(s, 9, 680, 262 + i * 170, 1140, 150, i + 1, col, lab, body, ic)
rrect(s, 680, 790, 1140, 72, C["ink"], radius=36)
text(s, 710, 790, 1080, 72, "Break-even items = Monthly costs ÷ Profit per item", 26,
     C["white"], font=BOLD_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
pill(s, 680, 880, 800, 76, C["lemon"], "Your ask: I need [amount] to [use of money].")
notes(s, NOTE_GEN + "\nThe chart is decoration only. Build your real forecast in a spreadsheet "
      "and ask an accountant or small-business advisor to review it.")

# ---------------- Slide 10: Ready to Launch! (centered, hero along the bottom) ----------------
s = new()
step_pill(s, 10)
ph(s, 10, "DECOR: aurora blob TL (Orchid > Bubblegum > Coral)", 0, 0, 300, 260)
ph(s, 10, "DECOR: aurora blob TR (Periwinkle > Aqua > Bubblegum)", 1620, 120, 300, 140)
ph(s, 10, "DECOR: sparkle confetti L", 30, 270, 240, 90)
ph(s, 10, "DECOR: sparkle confetti R", 1650, 270, 240, 90)
title(s, 260, 104, 1400, 146, "Ready to Launch!", size=120, align=PP_ALIGN.CENTER)
ph(s, 10, SQ, 700, 256, 520, 30)
text(s, 360, 298, 1200, 48, "Your plan is a living document — update it as you learn.", 28,
     C["ink"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
steps10 = [(C["orchid"], "Fill in each section"), (C["coral"], "Ask a friend to read it"),
           (C["peri"], "Test it with real customers"), (C["mint"], "Update it every 3 months")]
for i, (col, t) in enumerate(steps10):
    cx0 = 101 + i * 436
    rrect(s, cx0, 380, 410, 96, col, radius=48)
    ph(s, 10, "ICON: checkbox", cx0 + 18, 402, 52, 52)
    text(s, cx0 + 82, 380, [330, 262, 290, 240][i], 96, t, 26, on(col), font=BOLD_F, bold=True,
         anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
pills10 = [("123-456-7890", C["bubble"], 350, "phone"),
           ("hello@yourwebsite.com", C["peri"], 500, "mail"),
           ("www.yourwebsite.com", C["mint"], 480, "globe")]
x0 = (1920 - (sum(p[2] for p in pills10) + 2 * 36)) // 2
for lab, col, w, ic in pills10:
    pill(s, x0, 500, w, 76, col, lab, size=28, icon=ic, idx=10)
    x0 += w + 36
rrect(s, 260, 598, 1400, 64, tint(C["lemon"], 0.5), radius=32)
text(s, 290, 598, 1340, 64,
     "Based on common small-business planning guides. Adapt it to your own country and local "
     "rules.", 22, C["ink"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
ph(s, 10, "HERO: diverse makers cheering w/ flowers, coffee cups + a paper rocket, Idea Bulb "
          "buddy at the center 1400w", 260, 684, 1400, 330)

prs.save(OUT)
print("saved", OUT)
print("placeholders:", len(PLACEHOLDERS))
for row in PLACEHOLDERS:
    print("S%d | %s | %d,%d,%d,%d" % row)
